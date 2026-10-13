"""Unit tests for SlideCraft engine core: transitions, animations, geometry, effects, and icons."""

import os
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt

from slidecraft.engine.units import SLIDE_WIDTH_16_9, SLIDE_HEIGHT_16_9
from slidecraft.engine.presentation import SlideDeck
from slidecraft.engine.transitions import apply_morph_transition, apply_fade_transition, apply_push_transition
from slidecraft.engine.animations import AnimationBuilder
from slidecraft.engine.geometry import SVGPathToCustGeom, apply_svg_path_to_shape
from slidecraft.engine.effects import apply_soft_shadow, apply_gradient_fill, apply_solid_fill_with_opacity, set_corner_radius
from slidecraft.engine.icons import add_vector_icon, ICON_PATHS
from slidecraft.design_system.theme import get_theme

class TestSlideCraftEngine(unittest.TestCase):

    def setUp(self):
        self.output_dir = "test_output"
        os.makedirs(self.output_dir, exist_ok=True)

    def test_slide_deck_widescreen_dimensions(self):
        deck = SlideDeck()
        self.assertEqual(deck.prs.slide_width, SLIDE_WIDTH_16_9)
        self.assertEqual(deck.prs.slide_height, SLIDE_HEIGHT_16_9)

    def test_morph_transition_openxml(self):
        deck = SlideDeck()
        s1 = deck.add_slide()
        s2 = deck.add_slide(transition="morph", transition_duration_ms=1500)
        
        path = os.path.join(self.output_dir, "test_morph.pptx")
        deck.save(path)
        self.assertTrue(os.path.exists(path))

        # Inspect slide 2 XML
        with zipfile.ZipFile(path) as z:
            s2_xml = z.read("ppt/slides/slide2.xml").decode("utf-8")
            self.assertIn("p159:morph", s2_xml)
            self.assertIn('p14:dur="1500"', s2_xml)
            self.assertIn("AlternateContent", s2_xml)

    def test_animation_timing_openxml(self):
        deck = SlideDeck()
        slide = deck.add_slide()
        card = slide.add_card(Inches(1), Inches(1), Inches(4), Inches(3))
        slide.add_animation(card, effect="fade", delay_ms=300, dur_ms=500)
        
        path = os.path.join(self.output_dir, "test_animation.pptx")
        deck.save(path)
        self.assertTrue(os.path.exists(path))

        with zipfile.ZipFile(path) as z:
            s1_xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertIn("<p:timing", s1_xml)
            self.assertIn('filter="fade"', s1_xml)
            self.assertIn('delay="300"', s1_xml)

    def test_svg_path_to_custgeom(self):
        compiler = SVGPathToCustGeom()
        # Test cubic bezier and line path
        d = "M 0 0 L 10 0 C 15 5 15 15 10 20 L 0 20 Z"
        cust_geom = compiler.compile(d, view_box=(0, 0, 24, 24))
        self.assertIn("<a:custGeom", cust_geom)
        self.assertIn("<a:moveTo", cust_geom)
        self.assertIn("<a:lnTo", cust_geom)
        self.assertIn("<a:cubicBezTo", cust_geom)
        self.assertIn("<a:close/>", cust_geom)

        # Test elliptical arc to cubic Bezier conversion
        d_arc = "M 0 0 A 10 10 0 0 1 10 10 Z"
        cust_geom_arc = compiler.compile(d_arc, view_box=(0, 0, 24, 24))
        self.assertIn("<a:custGeom", cust_geom_arc)
        self.assertIn("<a:cubicBezTo", cust_geom_arc)

    def test_effects_shadow_and_gradient(self):
        deck = SlideDeck()
        slide = deck.add_slide()
        card = slide.add_card(Inches(1), Inches(1), Inches(4), Inches(3))
        apply_soft_shadow(card, blur_pt=16, opacity=0.25)
        apply_gradient_fill(card, "00F0FF", "8B5CF6", angle_deg=45)
        set_corner_radius(card, 7000)

        path = os.path.join(self.output_dir, "test_effects.pptx")
        deck.save(path)
        self.assertTrue(os.path.exists(path))

        with zipfile.ZipFile(path) as z:
            s1_xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertIn("<a:outerShdw", s1_xml)
            self.assertIn("<a:gradFill", s1_xml)
            self.assertIn('fmla="val 7000"', s1_xml)

    def test_vector_icons_generation(self):
        deck = SlideDeck()
        slide = deck.add_slide()
        # Add 5 diverse vector icons
        for idx, icon_name in enumerate(["shield", "brain", "zap", "database", "rocket"]):
            icon = slide.add_icon(icon_name, Inches(1 + idx * 2), Inches(2), size=Inches(0.8), bg_circle=True)
            self.assertIsNotNone(icon)

        path = os.path.join(self.output_dir, "test_icons.pptx")
        deck.save(path)
        self.assertTrue(os.path.exists(path))

        with zipfile.ZipFile(path) as z:
            s1_xml = z.read("ppt/slides/slide1.xml").decode("utf-8")
            self.assertIn("<a:custGeom", s1_xml)

    def test_embed_svg_vector_graphic(self):
        deck = SlideDeck()
        slide = deck.add_slide()
        svg_code = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100"><circle cx="50" cy="50" r="40" fill="#00F0FF"/></svg>'
        pic = slide.add_svg(svg_code, Inches(2), Inches(2), Inches(3), Inches(3))
        self.assertIsNotNone(pic)

        path = os.path.join(self.output_dir, "test_svg_embed.pptx")
        deck.save(path)
        self.assertTrue(os.path.exists(path))

        with zipfile.ZipFile(path) as z:
            self.assertTrue(any(f.startswith("ppt/media/") and f.endswith(".svg") for f in z.namelist()))

    def test_add_image_support(self):
        deck = SlideDeck()
        slide = deck.add_slide()
        img_path = os.path.abspath("images/sikkim_landscape.jpg")
        if os.path.exists(img_path):
            pic = slide.add_image(img_path, Inches(1), Inches(1), Inches(4), Inches(3), shadow=True, border_color="00F0FF")
            self.assertIsNotNone(pic)
            path = os.path.join(self.output_dir, "test_image.pptx")
            deck.save(path)
            self.assertTrue(os.path.exists(path))

    def test_powerpoint_desktop_com_openability(self):
        """Deep verification: confirm PowerPoint opens generated presentations without errors."""
        try:
            import win32com.client
            ppt = win32com.client.Dispatch("PowerPoint.Application")
        except Exception:
            self.skipTest("PowerPoint COM not available in this environment")

        test_files = [
            os.path.join(self.output_dir, "test_animation.pptx"),
            os.path.join(self.output_dir, "test_morph.pptx"),
            os.path.join(self.output_dir, "test_svg_embed.pptx")
        ]
        try:
            for f in test_files:
                if os.path.exists(f):
                    pres = ppt.Presentations.Open(os.path.abspath(f), ReadOnly=True, WithWindow=False)
                    self.assertGreater(pres.Slides.Count, 0)
                    pres.Close()
        finally:
            ppt.Quit()

if __name__ == "__main__":
    unittest.main()
