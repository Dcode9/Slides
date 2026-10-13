"""Unit tests for SlideCraft components: Bento grids, Cards, Metrics, Timelines, Comparisons."""

import os
import unittest
from pptx.util import Inches

from slidecraft.engine.presentation import SlideDeck
from slidecraft.components.hero import create_hero_title_slide
from slidecraft.components.header import add_slide_header
from slidecraft.components.cards import add_card_grid, add_feature_card
from slidecraft.components.bento import create_bento_grid
from slidecraft.components.metrics import add_metric_row, add_metric_card
from slidecraft.components.timeline import add_horizontal_timeline
from slidecraft.components.comparison import add_comparison_split
from slidecraft.components.quote import add_quote_callout
from slidecraft.design_system.theme import get_theme

class TestSlideCraftComponents(unittest.TestCase):

    def setUp(self):
        self.output_dir = "test_output"
        os.makedirs(self.output_dir, exist_ok=True)
        self.deck = SlideDeck(theme=get_theme("obsidian_cyber"))

    def test_hero_title_slide(self):
        slide = self.deck.add_slide()
        create_hero_title_slide(
            slide=slide,
            title="Next-Gen Architecture 2026",
            subtitle="Autonomous distributed computing.",
            eyebrow="EXECUTIVE BRIEFING",
            presenter="AI Systems Team",
            date="October 2026",
            accent_badge="CONFIDENTIAL"
        )
        self.assertGreater(len(slide.slide.shapes), 3)

    def test_card_grid_component(self):
        slide = self.deck.add_slide()
        add_slide_header(slide, title="Core Pillars", category="ARCHITECTURE")
        items = [
            {"title": "Zero-Trust Mesh", "body": "Continuous mutual attestation.", "icon": "shield"},
            {"title": "Neural Routing", "body": "Low-latency predictive packets.", "icon": "brain"},
            {"title": "Immutable Ledger", "body": "Cryptographic audit trails.", "icon": "database"}
        ]
        cards = add_card_grid(slide, items=items)
        self.assertEqual(len(cards), 3)

    def test_bento_grid_layouts(self):
        for layout in ["hero_left", "hero_right", "quad_asymmetric"]:
            slide = self.deck.add_slide()
            add_slide_header(slide, title=f"Bento Grid ({layout})", category="SYSTEM")
            items = [
                {"title": "Hero Card", "body": "Prominent capability overview.", "icon": "star", "type": "metric", "value": "10x", "label": "SPEEDUP"},
                {"title": "Feature B", "body": "High-throughput pipeline.", "icon": "zap"},
                {"title": "Feature C", "body": "Distributed telemetry.", "icon": "chart"},
                {"title": "Feature D", "body": "Fault tolerance.", "icon": "shield"}
            ]
            cards = create_bento_grid(slide, items=items, layout=layout)
            self.assertGreaterEqual(len(cards), 3)

    def test_metrics_row(self):
        slide = self.deck.add_slide()
        add_slide_header(slide, title="Performance Benchmarks", category="RESULTS")
        metrics = [
            {"value": "99.99%", "label": "Uptime", "delta": "+0.4%", "icon": "shield", "highlight": True},
            {"value": "1.2ms", "label": "Latency", "delta": "8x Faster", "icon": "zap"},
            {"value": "100M+", "label": "Requests", "delta": "+45% MoM", "icon": "trending_up"}
        ]
        cards = add_metric_row(slide, metrics=metrics)
        self.assertEqual(len(cards), 3)

    def test_horizontal_timeline(self):
        slide = self.deck.add_slide()
        add_slide_header(slide, title="Strategic Roadmap", category="TIMELINE")
        milestones = [
            {"date": "Q1 2026", "title": "Inception", "body": "Core architecture.", "icon": "lock"},
            {"date": "Q2 2026", "title": "Beta Mesh", "body": "Global testnet.", "icon": "globe"},
            {"date": "Q3 2026", "title": "Scale", "body": "Enterprise pilot.", "icon": "rocket"}
        ]
        cards = add_horizontal_timeline(slide, milestones=milestones)
        self.assertEqual(len(cards), 3)

    def test_comparison_split(self):
        slide = self.deck.add_slide()
        add_slide_header(slide, title="Legacy vs Modern", category="ANALYSIS")
        left_side = {
            "title": "Legacy Approach",
            "tag": "LEGACY",
            "points": ["Manual configurations", "Brittle deployments", "High latency"]
        }
        right_side = {
            "title": "Modern Platform",
            "tag": "NEXT-GEN",
            "points": ["Declarative policies", "Self-healing reconciliation", "Sub-millisecond routing"]
        }
        cards = add_comparison_split(slide, left_side=left_side, right_side=right_side)
        self.assertEqual(len(cards), 2)

    def test_quote_callout(self):
        slide = self.deck.add_slide()
        add_slide_header(slide, title="Customer Voice", category="TESTIMONIAL")
        card = add_quote_callout(
            slide=slide,
            quote_text="SlideCraft is completely redefining our presentation engineering workflow.",
            author="Jordan Matthews",
            author_title="Chief Technology Officer"
        )
        self.assertIsNotNone(card)

    def tearDown(self):
        output_file = os.path.join(self.output_dir, "test_all_components.pptx")
        self.deck.save(output_file)
        self.assertTrue(os.path.exists(output_file))

if __name__ == "__main__":
    unittest.main()
