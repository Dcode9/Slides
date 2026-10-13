"""Side-by-side comparison split component (Before vs After, Legacy vs Next-Gen)."""

from typing import List, Dict, Optional
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from slidecraft.engine.slide import SlideWrapper
from slidecraft.engine.units import DEFAULT_MARGIN_LEFT, CONTENT_WIDTH_16_9, hex_to_rgb
from slidecraft.engine.effects import set_corner_radius, apply_subtle_border

def add_comparison_split(
    slide: SlideWrapper,
    left_side: Dict,
    right_side: Dict,
    top=Inches(2.45),
    height=Inches(4.35),
    left=DEFAULT_MARGIN_LEFT,
    total_width=CONTENT_WIDTH_16_9,
    gap=Inches(0.4)
):
    """Render a side-by-side comparison split with subtle VS chip and feature checklists."""
    theme = slide.theme
    w_card = (total_width - gap) / 2

    # 1. Left Card (Legacy / Before)
    c_left = slide.add_card(
        left=left,
        top=top,
        width=w_card,
        height=height,
        bg_color=theme.palette.card_bg,
        border_color=theme.palette.card_border,
        border_width=1.0,
        morph_name="!!comp_left"
    )

    pad_x = Inches(0.4)
    y_l = top + Inches(0.35)

    # Left Tag
    slide.add_eyebrow_badge(
        text=left_side.get("tag", "TRADITIONAL / LEGACY"),
        left=left + pad_x,
        top=y_l,
        color_hex=theme.palette.text_muted,
        bg_hex=theme.palette.bg
    )
    y_l += Inches(0.45)

    # Left Title
    slide.add_text_box(
        text=left_side.get("title", ""),
        left=left + pad_x,
        top=y_l,
        width=w_card - (pad_x * 2),
        height=Inches(0.45),
        font_size=theme.scale.card_title,
        font_name=theme.fonts.heading_font,
        color_hex=theme.palette.text_primary,
        bold=True
    )
    y_l += Inches(0.55)

    # Left Bullets / Points
    for pt in left_side.get("points", []):
        slide.add_icon("cross_circle", left=left + pad_x, top=y_l, size=Inches(0.24), color_hex=theme.palette.danger)
        slide.add_text_box(
            text=pt,
            left=left + pad_x + Inches(0.35),
            top=y_l - Inches(0.04),
            width=w_card - (pad_x * 2) - Inches(0.35),
            height=Inches(0.35),
            font_size=theme.scale.body,
            font_name=theme.fonts.body_font,
            color_hex=theme.palette.text_muted
        )
        y_l += Inches(0.45)

    # 2. Right Card (Next-Gen / Solution) - Highlighted
    x_r = left + w_card + gap
    c_right = slide.add_card(
        left=x_r,
        top=top,
        width=w_card,
        height=height,
        bg_color=theme.palette.card_bg,
        border_color=theme.palette.accent,
        border_width=1.8,
        morph_name="!!comp_right"
    )

    y_r = top + Inches(0.35)

    # Right Tag
    slide.add_eyebrow_badge(
        text=right_side.get("tag", "NEXT-GEN / SLIDECRAFT"),
        left=x_r + pad_x,
        top=y_r,
        color_hex=theme.palette.accent,
        bg_hex=theme.palette.bg
    )
    y_r += Inches(0.45)

    # Right Title
    slide.add_text_box(
        text=right_side.get("title", ""),
        left=x_r + pad_x,
        top=y_r,
        width=w_card - (pad_x * 2),
        height=Inches(0.45),
        font_size=theme.scale.card_title,
        font_name=theme.fonts.heading_font,
        color_hex=theme.palette.text_primary,
        bold=True
    )
    y_r += Inches(0.55)

    # Right Bullets / Points
    for pt in right_side.get("points", []):
        slide.add_icon("check_circle", left=x_r + pad_x, top=y_r, size=Inches(0.24), color_hex=theme.palette.success)
        slide.add_text_box(
            text=pt,
            left=x_r + pad_x + Inches(0.35),
            top=y_r - Inches(0.04),
            width=w_card - (pad_x * 2) - Inches(0.35),
            height=Inches(0.35),
            font_size=theme.scale.body,
            font_name=theme.fonts.body_font,
            color_hex=theme.palette.text_primary
        )
        y_r += Inches(0.45)

    # 3. Central VS Badge
    vs_size = Inches(0.44)
    vs_x = left + w_card + (gap / 2) - (vs_size / 2)
    vs_y = top + (height / 2) - (vs_size / 2)
    vs_circle = slide.slide.shapes.add_shape(MSO_SHAPE.OVAL, vs_x, vs_y, vs_size, vs_size)
    vs_circle.fill.solid()
    vs_circle.fill.fore_color.rgb = hex_to_rgb(theme.palette.bg)
    apply_subtle_border(vs_circle, color_hex=theme.palette.card_border, width_pt=1.0)
    tf_vs = vs_circle.text_frame
    tf_vs.margin_top = Inches(0.04)
    p_vs = tf_vs.paragraphs[0]
    p_vs.text = "VS"
    p_vs.alignment = PP_ALIGN.CENTER
    p_vs.font.name = theme.fonts.heading_font
    p_vs.font.size = Pt(10)
    p_vs.font.bold = True
    p_vs.font.color.rgb = hex_to_rgb(theme.palette.accent)

    slide.stagger_shapes([c_left, c_right], effect="fly_in", direction="bottom", base_delay_ms=200, interval_ms=150)
    return [c_left, c_right]
