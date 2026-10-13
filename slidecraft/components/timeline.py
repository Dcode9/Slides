"""Process timeline and milestone roadmap components."""

from typing import List, Dict, Optional
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from slidecraft.engine.slide import SlideWrapper
from slidecraft.engine.units import DEFAULT_MARGIN_LEFT, CONTENT_WIDTH_16_9, hex_to_rgb
from slidecraft.components.base import calculate_columns
from slidecraft.components.cards import add_feature_card

def add_horizontal_timeline(
    slide: SlideWrapper,
    milestones: List[Dict],
    top=Inches(2.5),
    left=DEFAULT_MARGIN_LEFT,
    total_width=CONTENT_WIDTH_16_9,
    gap=Inches(0.3)
):
    """Add a horizontal roadmap/timeline with connected nodes and milestone cards."""
    theme = slide.theme
    count = len(milestones)
    cols = calculate_columns(total_width, count, gap=gap)

    # 1. Connecting Track Line
    track_y = top + Inches(0.2)
    track_line = slide.slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left + Inches(0.5),
        track_y,
        total_width - Inches(1.0),
        Inches(0.04)
    )
    track_line.fill.solid()
    track_line.fill.fore_color.rgb = hex_to_rgb(theme.palette.card_border)
    track_line.line.fill.background()

    card_shapes = []

    # 2. Render each milestone node and card
    for idx, (col_offset, col_w) in enumerate(cols):
        m = milestones[idx]
        node_center_x = left + col_offset + (col_w / 2)

        # Milestone Node Circle
        node_size = Inches(0.44)
        node = slide.slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            node_center_x - (node_size / 2),
            track_y - (node_size / 2) + Inches(0.02),
            node_size,
            node_size
        )
        node.fill.solid()
        node.fill.fore_color.rgb = hex_to_rgb(theme.palette.accent if idx == 0 else theme.palette.card_bg)
        node.line.color.rgb = hex_to_rgb(theme.palette.accent)
        node.line.width = Pt(1.5)

        # Node index number inside
        tf = node.text_frame
        tf.margin_top = Inches(0.02)
        p = tf.paragraphs[0]
        p.text = str(idx + 1)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = theme.fonts.heading_font
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(theme.palette.bg if idx == 0 else theme.palette.text_primary)

        # Milestone Card below
        card_top = top + Inches(0.65)
        card_h = Inches(3.65)
        card = add_feature_card(
            slide=slide,
            left=left + col_offset,
            top=card_top,
            width=col_w,
            height=card_h,
            title=m.get("title", ""),
            body=m.get("body", ""),
            icon=m.get("icon"),
            tag=m.get("date") or m.get("phase"),
            highlight=(idx == 0)
        )
        card_shapes.append(card)

    slide.stagger_shapes(card_shapes, effect="fly_in", direction="bottom", base_delay_ms=250, interval_ms=150)
    return card_shapes
