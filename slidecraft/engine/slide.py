"""High-level Slide wrapper providing layout math, surfaces, typography, and animation helpers."""

from typing import Optional, Union, List
from pptx.slide import Slide
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from slidecraft.engine.units import (
    SLIDE_WIDTH_16_9, SLIDE_HEIGHT_16_9,
    DEFAULT_MARGIN_LEFT, DEFAULT_MARGIN_TOP,
    hex_to_rgb, parse_color_hex
)
from slidecraft.engine.transitions import apply_transition, apply_morph_transition
from slidecraft.engine.animations import AnimationBuilder, AnimationEffect, AnimationTrigger
from slidecraft.engine.effects import (
    apply_soft_shadow, apply_gradient_fill, apply_subtle_border,
    apply_solid_fill_with_opacity, set_corner_radius
)
from slidecraft.engine.icons import add_vector_icon
from slidecraft.engine.svg_compiler import embed_svg_graphic
from slidecraft.design_system.theme import Theme

def set_shape_name(shape, name: str):
    """Set the internal OpenXML shape name (used for Morph matching with '!!' prefix)."""
    cNvPr = shape._element.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
    if cNvPr is not None:
        cNvPr.set('name', name)

class SlideWrapper:
    """Convenience wrapper around a PPTX Slide."""

    def __init__(self, slide: Slide, theme: Theme):
        self.slide = slide
        self.theme = theme
        self.animator = AnimationBuilder(slide)
        self.width = SLIDE_WIDTH_16_9
        self.height = SLIDE_HEIGHT_16_9

    def set_background_color(self, hex_color: Optional[str] = None):
        """Set solid slide background color."""
        color = hex_color or self.theme.palette.bg
        background = self.slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = hex_to_rgb(color)

    def set_morph(self, duration_ms: int = 1250, option: str = "byObject"):
        """Apply native Morph transition to this slide."""
        apply_morph_transition(self.slide, duration_ms=duration_ms, option=option)

    def set_transition(self, transition_type: str = "morph", duration_ms: int = 1000, **kwargs):
        """Apply slide transition."""
        apply_transition(self.slide, transition_type=transition_type, duration_ms=duration_ms, **kwargs)

    def add_card(
        self,
        left,
        top,
        width,
        height,
        bg_color: Optional[str] = None,
        border_color: Optional[str] = None,
        border_width: float = 1.0,
        radius: Optional[int] = None,
        shadow: Optional[bool] = None,
        morph_name: Optional[str] = None,
        opacity: float = 1.0
    ):
        """Create an elevated, modern rounded card surface with subtle border and drop shadow.

        Args:
            morph_name: If provided (e.g. '!!Card1'), enables Morph tracking across slides.
        """
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left,
            top,
            width,
            height
        )

        bg = bg_color or self.theme.palette.card_bg
        border = border_color or self.theme.palette.card_border
        r_val = radius if radius is not None else self.theme.card_radius
        has_shadow = shadow if shadow is not None else self.theme.card_shadow

        # Adjust corner radius
        set_corner_radius(shape, r_val)

        # Fill
        if opacity < 1.0:
            apply_solid_fill_with_opacity(shape, bg, opacity=opacity)
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = hex_to_rgb(bg)

        # Border
        if self.theme.card_border and border:
            apply_subtle_border(shape, color_hex=border, width_pt=border_width, opacity=self.theme.border_opacity)
        else:
            shape.line.fill.background()

        # Shadow
        if has_shadow:
            apply_soft_shadow(
                shape,
                blur_pt=self.theme.shadow_blur_pt,
                distance_pt=4.0,
                opacity=self.theme.shadow_opacity
            )

        if morph_name:
            set_shape_name(shape, morph_name)

        return shape

    def add_text_box(
        self,
        text: str,
        left,
        top,
        width,
        height,
        font_size: Optional[Pt] = None,
        font_name: Optional[str] = None,
        color_hex: Optional[str] = None,
        bold: bool = False,
        align: PP_ALIGN = PP_ALIGN.LEFT,
        line_spacing: float = 1.2,
        morph_name: Optional[str] = None
    ):
        """Add a cleanly styled text box with exact typographic hierarchy."""
        tx_box = self.slide.shapes.add_textbox(left, top, width, height)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0)
        tf.margin_right = Inches(0)
        tf.margin_top = Inches(0)
        tf.margin_bottom = Inches(0)

        p = tf.paragraphs[0]
        p.text = text
        p.alignment = align
        p.line_spacing = line_spacing

        p.font.size = font_size or self.theme.scale.body
        p.font.name = font_name or self.theme.fonts.body_font
        p.font.bold = bold
        p.font.color.rgb = hex_to_rgb(color_hex or self.theme.palette.text_primary)

        if morph_name:
            set_shape_name(tx_box, morph_name)

        return tx_box

    def add_eyebrow_badge(
        self,
        text: str,
        left,
        top,
        color_hex: Optional[str] = None,
        bg_hex: Optional[str] = None
    ):
        """Add a compact category/status pill tag (e.g. 'ARCHITECTURE • ENTERPRISE')."""
        color = color_hex or self.theme.palette.accent
        bg = bg_hex or self.theme.palette.card_bg

        # Add pill background
        pill_w = Inches(len(text) * 0.11 + 0.4)
        pill_h = Inches(0.3)
        pill = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left,
            top,
            pill_w,
            pill_h
        )
        set_corner_radius(pill, 25000) # Pill shape
        pill.fill.solid()
        pill.fill.fore_color.rgb = hex_to_rgb(bg)
        apply_subtle_border(pill, color_hex=color, width_pt=1.0, opacity=0.5)

        # Text
        tf = pill.text_frame
        tf.word_wrap = False
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.04)
        tf.margin_bottom = Inches(0.04)

        p = tf.paragraphs[0]
        p.text = text.upper()
        p.alignment = PP_ALIGN.CENTER
        p.font.name = self.theme.fonts.heading_font
        p.font.size = self.theme.scale.eyebrow
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(color)

        return pill

    def add_icon(
        self,
        icon_name: str,
        left,
        top,
        size=Inches(0.4),
        color_hex: Optional[str] = None,
        bg_circle: bool = False,
        bg_color_hex: Optional[str] = None
    ):
        """Add vector SVG icon to slide."""
        c = color_hex or self.theme.palette.accent
        bg = bg_color_hex or self.theme.palette.card_bg
        return add_vector_icon(
            self.slide,
            icon_name,
            left,
            top,
            size=size,
            color_hex=c,
            bg_circle=bg_circle,
            bg_color_hex=bg
        )

    def add_image(
        self,
        image_path: str,
        left,
        top,
        width,
        height,
        morph_name: Optional[str] = None,
        border_color: Optional[str] = None,
        border_width: float = 1.0,
        shadow: bool = False
    ):
        """Add an image picture shape to the slide with optional morph name, border, and shadow."""
        pic = self.slide.shapes.add_picture(image_path, left, top, width, height)
        if morph_name:
            set_shape_name(pic, morph_name)
        if border_color:
            apply_subtle_border(pic, color_hex=border_color, width_pt=border_width)
        if shadow:
            apply_soft_shadow(pic, blur_pt=self.theme.shadow_blur_pt, opacity=self.theme.shadow_opacity)
        return pic

    def add_svg(
        self,
        svg_content_or_path: str,
        left,
        top,
        width=Inches(2.0),
        height=Inches(2.0),
        name: str = "Vector Graphic"
    ):
        """Embed an SVG graphic natively into PowerPoint slide."""
        return embed_svg_graphic(self.slide, svg_content_or_path, left, top, width, height, name=name)

    def add_animation(
        self,
        shape,
        effect: AnimationEffect = "fade",
        trigger: AnimationTrigger = "withPrevious",
        delay_ms: int = 0,
        dur_ms: int = 500,
        direction: str = "bottom"
    ):
        """Add entrance animation to a shape."""
        self.animator.add_animation(
            shape=shape,
            effect=effect,
            trigger=trigger,
            delay_ms=delay_ms,
            dur_ms=dur_ms,
            direction=direction
        )

    def stagger_shapes(
        self,
        shapes: list,
        effect: AnimationEffect = "fade",
        base_delay_ms: int = 200,
        interval_ms: int = 150,
        dur_ms: int = 500,
        direction: str = "bottom"
    ):
        """Add staggered animations across multiple elements."""
        self.animator.stagger_shapes(
            shapes=shapes,
            effect=effect,
            base_delay_ms=base_delay_ms,
            interval_ms=interval_ms,
            dur_ms=dur_ms,
            direction=direction
        )

    def commit(self):
        """Finalize slide and compile timing tree."""
        self.animator.compile()
