"""SlideDeck: Top-level presentation orchestrator for SlideCraft."""

from typing import Optional, List
from pptx import Presentation
from pptx.slide import Slide

from slidecraft.engine.units import SLIDE_WIDTH_16_9, SLIDE_HEIGHT_16_9
from slidecraft.engine.slide import SlideWrapper
from slidecraft.design_system.theme import Theme, get_theme

class SlideDeck:
    """Manages creation, layout configuration, theme application, and compilation of PowerPoint decks."""

    def __init__(self, theme: Optional[Theme] = None, aspect_ratio: str = "16:9"):
        self.theme = theme or get_theme("obsidian_cyber")
        self.prs = Presentation()

        if aspect_ratio == "16:9":
            self.prs.slide_width = SLIDE_WIDTH_16_9
            self.prs.slide_height = SLIDE_HEIGHT_16_9

        self._blank_layout = self.prs.slide_layouts[6] # Blank slide layout
        self.slides: List[SlideWrapper] = []

    def add_slide(
        self,
        theme: Optional[Theme] = None,
        transition: str = "morph",
        transition_duration_ms: int = 1250,
        **transition_kwargs
    ) -> SlideWrapper:
        """Add a new slide to the deck with background color and transition pre-applied."""
        current_theme = theme or self.theme
        raw_slide = self.prs.slides.add_slide(self._blank_layout)
        wrapper = SlideWrapper(raw_slide, current_theme)

        # Set theme background color
        wrapper.set_background_color()

        # Set transition (first slide usually doesn't need morph, but can use fade)
        if len(self.slides) == 0 and transition == "morph":
            wrapper.set_transition("fade", duration_ms=800)
        else:
            wrapper.set_transition(transition, duration_ms=transition_duration_ms, **transition_kwargs)

        self.slides.append(wrapper)
        return wrapper

    def save(self, file_path: str):
        """Compile animations and save presentation to .pptx file."""
        for s in self.slides:
            s.commit()
        self.prs.save(file_path)
