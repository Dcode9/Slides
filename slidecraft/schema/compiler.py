import os
import json
from typing import Dict, Any, Union
from pptx.util import Inches
from slidecraft.engine.units import DEFAULT_MARGIN_LEFT, CONTENT_WIDTH_16_9
from slidecraft.engine.presentation import SlideDeck
from slidecraft.design_system.theme import get_theme
from slidecraft.components.hero import create_hero_title_slide
from slidecraft.components.header import add_slide_header
from slidecraft.components.cards import add_card_grid
from slidecraft.components.bento import create_bento_grid
from slidecraft.components.metrics import add_metric_row
from slidecraft.components.timeline import add_horizontal_timeline
from slidecraft.components.comparison import add_comparison_split
from slidecraft.components.quote import add_quote_callout

def compile_presentation_from_dict(spec: Dict[str, Any], output_path: str) -> SlideDeck:
    """Compile presentation specification dictionary into a .pptx file."""
    theme_data = spec.get("theme", {})
    palette_name = spec.get("palette") or theme_data.get("palette", "obsidian_cyber")
    fonts_name = spec.get("font_pairing") or theme_data.get("fonts", "modern_tech")
    aspect_ratio = spec.get("aspect_ratio", "16:9")

    theme = get_theme(palette_name=palette_name, font_name=fonts_name)
    deck = SlideDeck(theme=theme, aspect_ratio=aspect_ratio)

    slides_data = spec.get("slides", [])

    for idx, s_data in enumerate(slides_data):
        stype = s_data.get("type", "cards").lower()
        transition = s_data.get("transition", "morph" if idx > 0 else "fade")
        dur_ms = s_data.get("transition_duration_ms", 1250)

        slide = deck.add_slide(transition=transition, transition_duration_ms=dur_ms)

        if stype == "hero":
            create_hero_title_slide(
                slide=slide,
                title=s_data.get("title", ""),
                subtitle=s_data.get("subtitle"),
                eyebrow=s_data.get("eyebrow"),
                presenter=s_data.get("presenter"),
                date=s_data.get("date"),
                accent_badge=s_data.get("accent_badge"),
                center_aligned=s_data.get("center_aligned", False),
                image_path=s_data.get("image_path")
            )

        elif stype == "cards":
            add_slide_header(
                slide=slide,
                title=s_data.get("title", ""),
                category=s_data.get("category"),
                subtitle=s_data.get("subtitle")
            )
            items = s_data.get("items") or s_data.get("cards") or []
            if items:
                add_card_grid(slide=slide, items=items)

        elif stype == "bento":
            add_slide_header(
                slide=slide,
                title=s_data.get("title", ""),
                category=s_data.get("category"),
                subtitle=s_data.get("subtitle")
            )
            items = s_data.get("items") or s_data.get("cards") or []
            layout = s_data.get("layout", "hero_left")
            if items:
                create_bento_grid(slide=slide, items=items, layout=layout)

        elif stype == "metrics":
            add_slide_header(
                slide=slide,
                title=s_data.get("title", ""),
                category=s_data.get("category"),
                subtitle=s_data.get("subtitle")
            )
            items = s_data.get("items") or s_data.get("metrics") or []
            if items:
                add_metric_row(slide=slide, metrics=items)

        elif stype == "timeline":
            add_slide_header(
                slide=slide,
                title=s_data.get("title", ""),
                category=s_data.get("category"),
                subtitle=s_data.get("subtitle")
            )
            items = s_data.get("items") or s_data.get("steps") or s_data.get("milestones") or []
            if items:
                add_horizontal_timeline(slide=slide, milestones=items)

        elif stype == "comparison":
            add_slide_header(
                slide=slide,
                title=s_data.get("title", ""),
                category=s_data.get("category"),
                subtitle=s_data.get("subtitle")
            )
            left_side = s_data.get("left", {})
            right_side = s_data.get("right", {})
            add_comparison_split(slide=slide, left_side=left_side, right_side=right_side)

        elif stype == "quote":
            add_slide_header(
                slide=slide,
                title=s_data.get("title", ""),
                category=s_data.get("category"),
                subtitle=s_data.get("subtitle")
            )
            add_quote_callout(
                slide=slide,
                quote_text=s_data.get("quote", ""),
                author=s_data.get("author", ""),
                author_title=s_data.get("author_title")
            )

        elif stype in ("image", "gallery"):
            add_slide_header(
                slide=slide,
                title=s_data.get("title", ""),
                category=s_data.get("category"),
                subtitle=s_data.get("subtitle")
            )
            img_path = s_data.get("image_path")
            caption = s_data.get("caption") or s_data.get("body")
            if img_path and os.path.exists(img_path):
                if caption:
                    left_w = Inches(7.5)
                    img_h = Inches(4.3)
                    slide.add_image(
                        img_path,
                        left=DEFAULT_MARGIN_LEFT,
                        top=Inches(2.45),
                        width=left_w,
                        height=img_h,
                        border_color=slide.theme.palette.card_border,
                        shadow=True
                    )
                    rc_left = DEFAULT_MARGIN_LEFT + left_w + Inches(0.4)
                    rc_w = CONTENT_WIDTH_16_9 - left_w - Inches(0.4)
                    slide.add_card(
                        rc_left,
                        Inches(2.45),
                        rc_w,
                        img_h,
                        bg_color=slide.theme.palette.card_bg,
                        border_color=slide.theme.palette.card_border
                    )
                    slide.add_text_box(
                        text=caption,
                        left=rc_left + Inches(0.35),
                        top=Inches(2.8),
                        width=rc_w - Inches(0.7),
                        height=img_h - Inches(0.7),
                        font_size=slide.theme.scale.body,
                        font_name=slide.theme.fonts.body_font,
                        color_hex=slide.theme.palette.text_muted,
                        line_spacing=1.3
                    )
                else:
                    slide.add_image(
                        img_path,
                        left=DEFAULT_MARGIN_LEFT + Inches(1.0),
                        top=Inches(2.45),
                        width=CONTENT_WIDTH_16_9 - Inches(2.0),
                        height=Inches(4.3),
                        border_color=slide.theme.palette.card_border,
                        shadow=True
                    )

    deck.save(output_path)
    return deck

def compile_presentation_from_json(json_content_or_path: str, output_path: str) -> SlideDeck:
    """Load JSON from string or file path and compile into .pptx."""
    import os
    if os.path.exists(json_content_or_path):
        with open(json_content_or_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = json.loads(json_content_or_path)

    return compile_presentation_from_dict(data, output_path)
